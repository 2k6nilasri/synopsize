import pandas as pd
import io
from typing import Dict, Any, List
from app.core.security import neutralize_formula_injection

def process_excel_or_csv(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Parses Excel (.xlsx) or CSV (.csv) spreadsheets.
    Reads dataframes, neutralizes formula injection, and builds structured markdown table blocks.
    """
    ext = filename.split(".")[-1].lower()
    sheets_data = {}
    sheets_metadata = {}
    
    if ext == "csv":
        df = pd.read_csv(io.BytesIO(file_bytes))
        sheets_data["Sheet1"] = df
    else:
        excel_file = pd.ExcelFile(io.BytesIO(file_bytes))
        for sheet in excel_file.sheet_names:
            sheets_data[sheet] = excel_file.parse(sheet)
        if ext == "xlsx":
            from openpyxl import load_workbook

            formulas_book = load_workbook(io.BytesIO(file_bytes), data_only=False, read_only=True)
            values_book = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
            for worksheet in formulas_book.worksheets:
                values_sheet = values_book[worksheet.title]
                formulas = []
                for row in worksheet.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and cell.value.startswith("="):
                            cached_value = values_sheet[cell.coordinate].value
                            if cached_value is not None and not isinstance(
                                cached_value, (str, int, float, bool)
                            ):
                                cached_value = (
                                    cached_value.isoformat()
                                    if hasattr(cached_value, "isoformat")
                                    else str(cached_value)
                                )
                            formulas.append(
                                {
                                    "cell": cell.coordinate,
                                    "formula": cell.value,
                                    "cached_value": cached_value,
                                }
                            )
                sheets_metadata[worksheet.title] = {
                    "visibility": worksheet.sheet_state,
                    "formulas": formulas,
                }
            formulas_book.close()
            values_book.close()

    pages = []
    ocr_lines = []
    
    for page_idx, (sheet_name, df) in enumerate(sheets_data.items()):
        page_num = page_idx + 1
        blocks = []
        
        # Clean dataframe strings & neutralize formula injection
        df = df.fillna("")
        columns = [neutralize_formula_injection(str(c)) for c in df.columns]
        
        # Heading block
        blocks.append({
            "id": f"block-{page_num}-1",
            "type": "heading",
            "reading_order": 1,
            "page": page_num,
            "bbox": [40, 40, 750, 80],
            "content": f"Sheet: {sheet_name}",
            "confidence": 0.99,
            "extractor": "Pandas Excel/CSV Engine",
            "source_reference": f"spreadsheet_sheet_{sheet_name}_heading"
        })
        
        # Table block
        header_row = "| " + " | ".join(columns) + " |"
        separator_row = "| " + " | ".join(["---"] * len(columns)) + " |"
        data_rows = []
        
        for r_idx, row in df.iterrows():
            clean_vals = [neutralize_formula_injection(str(v).replace("\n", " ")) for v in row.values]
            data_rows.append("| " + " | ".join(clean_vals) + " |")
            
        md_table = "\n".join([header_row, separator_row] + data_rows)
        
        blocks.append({
            "id": f"block-{page_num}-2",
            "type": "table",
            "reading_order": 2,
            "page": page_num,
            "bbox": [40, 100, 750, 600],
            "content": md_table,
            "confidence": 0.98,
            "extractor": "Pandas Excel/CSV Engine",
            "source_reference": f"spreadsheet_sheet_{sheet_name}_table",
            "metadata": {
                "sheet_name": sheet_name,
                "total_rows": len(df),
                "total_cols": len(columns),
                "columns": columns,
                **sheets_metadata.get(sheet_name, {"visibility": "visible", "formulas": []}),
            }
        })
        
        ocr_lines.append(f"--- Sheet {sheet_name} ---")
        ocr_lines.append(f"Header [0.99]: {', '.join(columns)}")
        ocr_lines.append(f"Rows count: {len(df)}")
        
        page_img_url = f"data:image/svg+xml;utf8,<svg xmlns='http://www.w3.org/2000/svg' width='800' height='700' viewBox='0 0 800 700'><rect width='100%' height='100%' fill='%23f8fafc'/><text x='50%' y='50%' dominant-baseline='middle' text-anchor='middle' font-family='sans-serif' font-size='20' fill='%23475569'>Spreadsheet Render: {sheet_name} ({len(df)} rows)</text></svg>"

        pages.append({
            "page_number": page_num,
            "width": 800,
            "height": 700,
            "image_data": page_img_url,
            "blocks": blocks
        })

    return {
        "total_pages": len(pages),
        "pages": pages,
        "ocr_text": "\n".join(ocr_lines)
    }
