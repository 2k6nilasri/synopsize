import base64
import io
from typing import Any, Dict, List

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def _preview(slide_number: int) -> str:
    svg = (
        '<svg xmlns="http://www.w3.org/2000/svg" width="1280" height="720">'
        '<rect width="100%" height="100%" fill="#f8fafc"/>'
        f'<text x="50%" y="50%" text-anchor="middle" font-family="sans-serif" '
        f'font-size="28" fill="#475569">Presentation slide {slide_number}</text></svg>'
    )
    return "data:image/svg+xml;base64," + base64.b64encode(svg.encode()).decode()


def _table_markdown(rows: List[List[str]]) -> str:
    width = max((len(row) for row in rows), default=0)
    if not width:
        return ""
    normalized = [row + [""] * (width - len(row)) for row in rows]
    lines = [normalized[0], ["---"] * width, *normalized[1:]]
    return "\n".join(
        "| " + " | ".join(value.replace("|", "\\|") for value in row) + " |"
        for row in lines
    )


def _shape_bbox(shape: Any) -> List[int]:
    scale = 96 / 914400
    return [
        int(shape.left * scale),
        int(shape.top * scale),
        int((shape.left + shape.width) * scale),
        int((shape.top + shape.height) * scale),
    ]


def process_pptx_document(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    presentation = Presentation(io.BytesIO(file_bytes))
    pages: List[Dict[str, Any]] = []
    ocr_lines: List[str] = []
    slide_width = int(presentation.slide_width / 914400 * 96)
    slide_height = int(presentation.slide_height / 914400 * 96)

    for slide_number, slide in enumerate(presentation.slides, start=1):
        blocks: List[Dict[str, Any]] = []
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                continue
            text = shape.text.strip() if shape.has_text_frame else ""
            if shape.has_table:
                rows = [
                    [cell.text.strip() for cell in row.cells]
                    for row in shape.table.rows
                ]
                if rows:
                    blocks.append(
                        {
                            "type": "table",
                            "content": _table_markdown(rows),
                            "metadata": {"raw_table": rows},
                            "confidence": 0.98,
                            "bbox": _shape_bbox(shape),
                        }
                    )
            elif text:
                is_title = shape == slide.shapes.title
                blocks.append(
                    {
                        "type": "heading" if is_title else "paragraph",
                        "content": text,
                        "confidence": 0.98,
                        "bbox": _shape_bbox(shape),
                    }
                )
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = shape.image
                image_data = (
                    "data:"
                    + image.content_type
                    + ";base64,"
                    + base64.b64encode(image.blob).decode("ascii")
                )
                blocks.append(
                    {
                        "type": "figure",
                        "content": f"![{shape.name}]({image_data})",
                        "metadata": {"image_data": image_data, "name": shape.name},
                        "confidence": 0.8,
                        "bbox": _shape_bbox(shape),
                    }
                )

            if shape.has_chart:
                chart = shape.chart
                chart_data = {
                    "chart_type": str(chart.chart_type),
                    "series": [
                        {
                            "name": series.name,
                            "values": list(series.values),
                        }
                        for series in chart.series
                    ],
                }
                blocks.append(
                    {
                        "type": "chart",
                        "content": f"Chart: {chart.chart_title.text_frame.text if chart.has_title else shape.name}",
                        "metadata": chart_data,
                        "confidence": 0.9,
                        "bbox": _shape_bbox(shape),
                    }
                )

        try:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        except (AttributeError, KeyError, ValueError):
            notes = ""
        if notes:
            blocks.append(
                {
                    "type": "paragraph",
                    "content": f"Speaker notes: {notes}",
                    "confidence": 0.98,
                    "extractor": "python-pptx speaker notes",
                }
            )

        for reading_order, block in enumerate(blocks, start=1):
            block.setdefault("id", f"block-{slide_number}-{reading_order}")
            block.setdefault("reading_order", reading_order)
            block.setdefault("page", slide_number)
            block.setdefault("bbox", [0, (reading_order - 1) * 28, slide_width, reading_order * 28])
            block.setdefault("confidence", 0.95)
            block.setdefault("extractor", "python-pptx")
            block.setdefault("source_reference", f"pptx_slide_{slide_number}_{reading_order}")
            ocr_lines.append(
                f"Slide {slide_number}, block {reading_order} "
                f"[{block['confidence']:.2f}]: {block['content']}"
            )

        pages.append(
            {
                "page_number": slide_number,
                "width": slide_width,
                "height": slide_height,
                "image_data": _preview(slide_number),
                "blocks": blocks,
            }
        )

    return {
        "total_pages": len(pages),
        "pages": pages,
        "ocr_text": "\n".join(ocr_lines),
    }
