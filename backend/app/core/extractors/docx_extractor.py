import docx
import io
import base64
import zipfile
from xml.etree import ElementTree
from typing import Dict, Any, List

def process_docx_document(docx_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Parses Microsoft Word (.docx) documents.
    Extracts headings, paragraphs, structured markdown tables, and inline images.
    """
    doc = docx.Document(io.BytesIO(docx_bytes))
    blocks = []
    reading_order = 1
    ocr_lines = ["--- Page 1 (DOCX Document) ---"]
    revisions = []
    comments = []
    footnotes = []
    endnotes = []
    
    # Process Paragraphs & Headings
    for p in doc.paragraphs:
        accepted_text = "".join(p._p.xpath(".//w:t/text()"))
        text = (accepted_text or p.text).strip()
        if not text:
            continue
            
        style_name = p.style.name.lower() if p.style else ""
        if "heading" in style_name or style_name.startswith("h1") or style_name.startswith("h2"):
            b_type = "heading"
            conf = 0.99
        elif "list" in style_name or p.style.name.startswith("List"):
            b_type = "paragraph"
            text = f"- {text}"
            conf = 0.95
        else:
            b_type = "paragraph"
            conf = 0.96

        blocks.append({
            "id": f"block-1-{reading_order}",
            "type": b_type,
            "reading_order": reading_order,
            "page": 1,
            "bbox": [50, 40 + reading_order * 45, 750, 80 + reading_order * 45],
            "content": text,
            "confidence": conf,
            "extractor": "python-docx Parser",
            "source_reference": f"docx_paragraph_{reading_order}"
        })
        ocr_lines.append(f"Line {reading_order} [{conf:.2f}]: {text}")
        reading_order += 1

    for revision in doc._element.body.xpath(".//w:ins | .//w:del"):
        deleted_text = "".join(revision.xpath(".//w:delText/text()"))
        inserted_text = "".join(revision.xpath(".//w:t/text()"))
        revisions.append(
            {
                "author": revision.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}author"),
                "date": revision.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}date"),
                "kind": "deleted" if revision.tag.endswith("}del") else "inserted",
                "text": deleted_text or inserted_text,
            }
        )

    # Process Tables
    for t_idx, table in enumerate(doc.tables):
        if not table.rows:
            continue
            
        table_rows = []
        for row in table.rows:
            row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            table_rows.append(row_cells)

        if table_rows:
            headers = table_rows[0]
            data_rows = table_rows[1:]
            md_table = "| " + " | ".join(headers) + " |\n"
            md_table += "| " + " | ".join(["---"] * len(headers)) + " |\n"
            for r in data_rows:
                md_table += "| " + " | ".join(r) + " |\n"

            blocks.append({
                "id": f"block-1-{reading_order}",
                "type": "table",
                "reading_order": reading_order,
                "page": 1,
                "bbox": [50, 40 + reading_order * 55, 750, 180 + reading_order * 55],
                "content": md_table,
                "confidence": 0.97,
                "extractor": "python-docx Table Parser",
                "source_reference": f"docx_table_{t_idx+1}",
                "metadata": {
                    "raw_table": table_rows,
                    "rows": len(data_rows),
                    "columns": len(headers)
                }
            })
            ocr_lines.append(f"Table {t_idx+1} [0.97]: {len(data_rows)} rows x {len(headers)} cols")
            reading_order += 1

    for section_index, section in enumerate(doc.sections, start=1):
        for label, container in (("Header", section.header), ("Footer", section.footer)):
            text = "\n".join(
                paragraph.text.strip()
                for paragraph in container.paragraphs
                if paragraph.text.strip()
            )
            if text:
                blocks.append(
                    {
                        "id": f"block-1-{reading_order}",
                        "type": "header_footer",
                        "reading_order": reading_order,
                        "page": 1,
                        "bbox": [0, 0, 800, 40],
                        "content": text,
                        "confidence": 0.98,
                        "extractor": "python-docx header/footer parser",
                        "source_reference": f"docx_{label.lower()}_{section_index}",
                    }
                )
                reading_order += 1

    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as archive:
        names = set(archive.namelist())
        if "word/comments.xml" in names:
            root = ElementTree.fromstring(archive.read("word/comments.xml"))
            comment_ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            for comment in root.findall(f"{comment_ns}comment"):
                comments.append(
                    {
                        "id": comment.get(f"{comment_ns}id"),
                        "author": comment.get(f"{comment_ns}author"),
                        "date": comment.get(f"{comment_ns}date"),
                        "text": "".join(
                            element.text or ""
                            for element in comment.iter(f"{comment_ns}t")
                        ),
                    }
                )
        for part_name, target in (
            ("word/footnotes.xml", footnotes),
            ("word/endnotes.xml", endnotes),
        ):
            if part_name not in names:
                continue
            root = ElementTree.fromstring(archive.read(part_name))
            namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            for note in root:
                note_id = note.get(f"{namespace}id")
                if note_id in {"-1", "0"}:
                    continue
                note_text = "".join(
                    element.text or "" for element in note.iter(f"{namespace}t")
                )
                if note_text:
                    target.append({"id": note_id, "text": note_text})

    for note_type, notes in (("footnote", footnotes), ("endnote", endnotes)):
        for note in notes:
            blocks.append(
                {
                    "id": f"block-1-{reading_order}",
                    "type": "footnote",
                    "reading_order": reading_order,
                    "page": 1,
                    "bbox": [40, 900, 760, 980],
                    "content": note["text"],
                    "confidence": 0.96,
                    "extractor": "python-docx note parser",
                    "source_reference": f"docx_{note_type}_{note['id']}",
                }
            )
            reading_order += 1

    for image_index, part in enumerate(doc.part.related_parts.values(), start=1):
        content_type = getattr(part, "content_type", "")
        if not content_type.startswith("image/"):
            continue
        image_data = (
            f"data:{content_type};base64,"
            + base64.b64encode(part.blob).decode("ascii")
        )
        blocks.append(
            {
                "id": f"block-1-{reading_order}",
                "type": "figure",
                "reading_order": reading_order,
                "page": 1,
                "bbox": [50, 40 + reading_order * 45, 750, 280 + reading_order * 45],
                "content": f"![Embedded Word image {image_index}]({image_data})",
                "confidence": 0.9,
                "extractor": "python-docx image parser",
                "source_reference": f"docx_image_{image_index}",
                "metadata": {"image_data": image_data},
            }
        )
        reading_order += 1

    # Default placeholder image preview for DOCX
    page_img_url = "data:image/svg+xml;utf8,<svg xmlns='http://www.w3.org/2000/svg' width='800' height='1000' viewBox='0 0 800 1000'><rect width='100%' height='100%' fill='%23f8fafc'/><text x='50%' y='50%' dominant-baseline='middle' text-anchor='middle' font-family='sans-serif' font-size='20' fill='%2364748b'>Word Document (.docx) Render</text></svg>"

    return {
        "total_pages": 1,
        "pages": [
            {
                "page_number": 1,
                "width": 800,
                "height": 1000,
                "image_data": page_img_url,
                "blocks": blocks
            }
        ],
        "ocr_text": "\n".join(ocr_lines),
        "revisions": revisions,
        "comments": comments,
        "footnotes": footnotes,
        "endnotes": endnotes,
    }
