import io
import json
import os
from email.message import EmailMessage

from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
import pymupdf as fitz

from app.cli import parse_file
from app.main import app
import app.api.endpoints as api_endpoints
import app.core.audit as audit_store
from app.core.file_parser import extract_document
from app.core.extractors.image_extractor import process_image_file
from app.core.security import validate_uploaded_file
from app.core.audit import append_audit_event, verify_audit_log
from app.core.retention import cleanup_expired_artifacts

client = TestClient(app)


def test_text_and_html_extractors_preserve_structure_and_remove_scripts():
    html_bytes = (
        b"<html><body><h1>Quarterly report</h1><script>secret()</script>"
        b"<p>Revenue increased.</p><table><tr><th>Year</th><th>Value</th></tr>"
        b"<tr><td>2026</td><td>12</td></tr></table></body></html>"
    )
    parsed = extract_document(html_bytes, "report.html")
    blocks = parsed["pages"][0]["blocks"]

    assert [block["type"] for block in blocks] == ["heading", "paragraph", "table"]
    assert all("secret" not in block["content"] for block in blocks)
    assert "Revenue increased." in parsed["ocr_text"]
    assert validate_uploaded_file(html_bytes, "report.html", store=False)["is_valid"]


def test_eml_attachment_is_parsed_as_child_document():
    message = EmailMessage()
    message["From"] = "analyst@example.com"
    message["To"] = "team@example.com"
    message["Subject"] = "Review"
    message.set_content("Please review the attached notes.")
    message.add_attachment(
        b"# Notes\n\nCheck the totals.",
        maintype="text",
        subtype="markdown",
        filename="notes.md",
    )

    parsed = extract_document(message.as_bytes(), "review.eml")
    child = parsed["child_documents"][0]
    assert child["filename"] == "notes.md"
    assert child["pages"][0]["blocks"][0]["type"] == "heading"
    assert "Subject: Review" in parsed["ocr_text"]


def test_nested_eml_attachment_is_recursively_parsed():
    nested = EmailMessage()
    nested["Subject"] = "Forwarded details"
    nested.set_content("Nested email body.")
    outer = EmailMessage()
    outer["Subject"] = "Outer message"
    outer.set_content("See forwarded email.")
    outer.add_attachment(
        nested.as_bytes(),
        maintype="message",
        subtype="rfc822",
        filename="forwarded.eml",
    )

    parsed = extract_document(outer.as_bytes(), "outer.eml")

    assert parsed["child_documents"][0]["filename"] == "forwarded.eml"
    assert any(
        block["content"] == "Nested email body."
        for block in parsed["child_documents"][0]["pages"][0]["blocks"]
    )


def test_pptx_extractor_includes_slide_text_tables_and_speaker_notes():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Operational review"
    table_shape = slide.shapes.add_table(2, 2, 0, 0, 3_000_000, 1_000_000)
    table_shape.table.cell(0, 0).text = "Metric"
    table_shape.table.cell(0, 1).text = "Value"
    table_shape.table.cell(1, 0).text = "Revenue"
    table_shape.table.cell(1, 1).text = "120"
    slide.notes_slide.notes_text_frame.text = "Discuss the revenue variance."
    buffer = io.BytesIO()
    presentation.save(buffer)

    parsed = extract_document(buffer.getvalue(), "review.pptx")
    blocks = parsed["pages"][0]["blocks"]
    assert any(block["type"] == "heading" and block["content"] == "Operational review" for block in blocks)
    assert any(block["type"] == "table" for block in blocks)
    assert any("Speaker notes:" in block["content"] for block in blocks)
    assert validate_uploaded_file(buffer.getvalue(), "review.pptx", store=False)["is_valid"]


def test_pdf_short_equation_is_extracted_as_equation_with_latex():
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 120), "x^2 + y^2 = z^2", fontsize=14)
    pdf_bytes = document.tobytes()
    document.close()

    parsed = extract_document(pdf_bytes, "equation.pdf")
    equation = next(
        block
        for block in parsed["pages"][0]["blocks"]
        if block["type"] == "equation"
    )

    assert equation["metadata"]["latex"] == equation["content"]
    assert equation["content"].startswith("\\[")
    assert "x^2 + y^2 = z^2" in equation["content"]


def test_multipage_tiff_and_xlsx_metadata_are_preserved():
    first = Image.new("RGB", (80, 80), "white")
    second = Image.new("RGB", (80, 80), "white")
    image_buffer = io.BytesIO()
    first.save(image_buffer, format="TIFF", save_all=True, append_images=[second])
    image_result = extract_document(image_buffer.getvalue(), "scan.tiff")
    assert image_result["total_pages"] == 2
    assert [page["page_number"] for page in image_result["pages"]] == [1, 2]

    workbook = Workbook()
    visible = workbook.active
    visible.title = "Summary"
    visible.append(["Total"])
    visible.append(["=SUM(1,2)"])
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["Internal"])
    workbook_buffer = io.BytesIO()
    workbook.save(workbook_buffer)

    spreadsheet = extract_document(workbook_buffer.getvalue(), "metrics.xlsx")
    summary = next(
        block
        for block in spreadsheet["pages"][0]["blocks"]
        if block["type"] == "table"
    )
    hidden_page = spreadsheet["pages"][1]
    hidden_table = next(
        block for block in hidden_page["blocks"] if block["type"] == "table"
    )
    assert summary["metadata"]["formulas"][0]["formula"] == "=SUM(1,2)"
    assert hidden_table["metadata"]["visibility"] == "hidden"


def test_image_ocr_failure_is_flagged_without_fabricated_text(monkeypatch):
    image_buffer = io.BytesIO()
    Image.new("RGB", (80, 80), "white").save(image_buffer, format="PNG")

    def fail_ocr(*args, **kwargs):
        raise RuntimeError("OCR engine unavailable")

    monkeypatch.setattr(
        "app.core.extractors.image_extractor.pytesseract.image_to_data",
        fail_ocr,
    )
    parsed = process_image_file(image_buffer.getvalue(), "unreadable.png")
    block = parsed["pages"][0]["blocks"][0]

    assert block["content"] == ""
    assert block["confidence"] == 0.0
    assert "needs_review" in block["flags"]
    assert "OCR engine unavailable" in parsed["ocr_text"]


def test_cli_writes_structured_outputs(monkeypatch, tmp_path):
    monkeypatch.setattr(audit_store, "AUDIT_LOG_FILE", tmp_path / "audit.jsonl")
    source = tmp_path / "summary.txt"
    source.write_text("A short business summary.", encoding="utf-8")
    output_dir = tmp_path / "parsed"

    result_code = parse_file(source, output_dir, no_embeddings=True)
    result = json.loads((output_dir / "result.json").read_text(encoding="utf-8"))

    assert result_code == 0
    assert result["pages"][0]["blocks"][0]["provenance"]["file_sha256"]
    assert (output_dir / "result.md").is_file()
    assert (output_dir / "ocr.txt").is_file()
    assert (output_dir / "chunks.jsonl").is_file()
    assert (output_dir / "report.json").is_file()
    assert (output_dir / "figures").is_dir()


def test_api_parse_returns_provenance_schema_for_text_documents(monkeypatch, tmp_path):
    monkeypatch.setattr(audit_store, "AUDIT_LOG_FILE", tmp_path / "audit.jsonl")
    response = client.post(
        "/api/parse",
        files={"file": ("brief.md", b"# Overview\nThe service is ready.", "text/markdown")},
    )
    assert response.status_code == 200
    job_id = response.json()["job_id"]
    try:
        result_response = client.get(f"/api/jobs/{job_id}/result")
        assert result_response.status_code == 200
        block = result_response.json()["pages"][0]["blocks"][0]
        assert block["type"] == "heading"
        assert len(block["bbox"]) == 4
        assert block["provenance"]["file_sha256"]
        assert block["routing"]
        assert result_response.json()["audit_entry_hash"]
        audit_response = client.get("/api/audit")
        assert audit_response.json()["integrity_valid"] is True
        assert audit_response.json()["verified_entries"] == 1
    finally:
        client.delete(f"/api/jobs/{job_id}")


def test_api_uploads_pptx_and_returns_slide_content(monkeypatch, tmp_path):
    monkeypatch.setattr(audit_store, "AUDIT_LOG_FILE", tmp_path / "audit.jsonl")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Quarterly results"
    textbox = slide.shapes.add_textbox(0, 1_000_000, 5_000_000, 500_000)
    textbox.text_frame.text = "Revenue increased by 12 percent."
    slide.notes_slide.notes_text_frame.text = "Discuss the growth drivers."
    buffer = io.BytesIO()
    presentation.save(buffer)

    response = client.post(
        "/api/parse",
        files={
            "file": (
                "quarterly-results.pptx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )
        },
    )
    assert response.status_code == 200
    job_id = response.json()["job_id"]
    try:
        result_response = client.get(f"/api/jobs/{job_id}/result")
        assert result_response.status_code == 200
        result = result_response.json()
        assert result["total_pages"] == 1
        blocks = result["pages"][0]["blocks"]
        assert any(block["content"] == "Quarterly results" for block in blocks)
        assert any(
            block["content"] == "Revenue increased by 12 percent."
            for block in blocks
        )
        assert any(
            "Speaker notes: Discuss the growth drivers." in block["content"]
            for block in blocks
        )
        assert "Revenue increased by 12 percent." in result["markdown"]
    finally:
        client.delete(f"/api/jobs/{job_id}")


def test_api_returns_structured_error_for_unsupported_format():
    response = client.post(
        "/api/parse",
        files={"file": ("payload.exe", b"not executable data", "application/octet-stream")},
    )

    assert response.status_code == 400
    detail = response.json()["detail"]
    assert detail["status"] == "error"
    assert detail["stage"] == 2
    assert detail["code"] == "UNSUPPORTED_FORMAT"


def test_api_limits_image_tool_uploads_and_rejects_unknown_job_ids(monkeypatch):
    monkeypatch.setattr(api_endpoints, "MAX_FILE_SIZE_BYTES", 8)
    oversized_response = client.post(
        "/api/tools/denoise",
        files={"file": ("scan.png", b"\x89PNG\r\n\x1a\npayload", "image/png")},
    )
    assert oversized_response.status_code == 413

    assert client.get("/api/jobs/not-a-uuid/result").status_code == 400
    missing_job = "00000000-0000-0000-0000-000000000000"
    assert client.get(f"/api/jobs/{missing_job}/events").status_code == 404
    assert client.delete(f"/api/jobs/{missing_job}").status_code == 404


def test_cors_allows_configured_ui_origin_only():
    allowed = client.options(
        "/api/validate",
        headers={
            "Origin": "http://localhost:3000",
            "Access-Control-Request-Method": "POST",
            "Access-Control-Request-Headers": "content-type",
        },
    )
    allowed_fallback_port = client.options(
        "/api/validate",
        headers={
            "Origin": "http://localhost:3002",
            "Access-Control-Request-Method": "POST",
            "Access-Control-Request-Headers": "content-type",
        },
    )
    rejected = client.options(
        "/api/validate",
        headers={
            "Origin": "https://untrusted.example",
            "Access-Control-Request-Method": "POST",
            "Access-Control-Request-Headers": "content-type",
        },
    )

    assert allowed.headers["access-control-allow-origin"] == "http://localhost:3000"
    assert (
        allowed_fallback_port.headers["access-control-allow-origin"]
        == "http://localhost:3002"
    )
    assert "access-control-allow-origin" not in rejected.headers


def test_audit_chain_detects_tampering(monkeypatch, tmp_path):
    audit_path = tmp_path / "audit.jsonl"
    monkeypatch.setattr(audit_store, "AUDIT_LOG_FILE", audit_path)
    append_audit_event("job-a", "document_parse", "completed", "abc123", stage=7)
    append_audit_event("job-b", "document_parse", "error", "def456", stage=4)
    assert verify_audit_log() == (True, 2)

    records = audit_path.read_text(encoding="utf-8").splitlines()
    records[0] = records[0].replace("abc123", "tampered")
    audit_path.write_text("\n".join(records) + "\n", encoding="utf-8")
    assert verify_audit_log() == (False, 1)


def test_retention_removes_only_expired_uploads(tmp_path):
    expired = tmp_path / "expired.json"
    retained = tmp_path / "retained.json"
    expired.write_text("old", encoding="utf-8")
    retained.write_text("new", encoding="utf-8")
    expired_upload_dir = tmp_path / "expired-upload"
    expired_upload_dir.mkdir()
    expired_upload = expired_upload_dir / "file.pdf"
    expired_upload.write_bytes(b"old")
    os.utime(expired, (0, 0))
    os.utime(expired_upload, (0, 0))
    os.utime(retained, (90 * 3600, 90 * 3600))

    removed = cleanup_expired_artifacts(
        now=100 * 3600,
        retention_hours=24,
        storage_directories=[tmp_path],
    )

    assert removed == 2
    assert not expired.exists()
    assert not expired_upload_dir.exists()
    assert retained.is_file()
